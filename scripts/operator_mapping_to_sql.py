"""Convert the curated operator mapping sheet into a dbmate migration.

Usage: py -3 scripts/operator_mapping_to_sql.py db/mappings/operator-mapping.xlsx \
           db/migrations/20261001000123_operator_backfill.sql
"""
import sys

from openpyxl import load_workbook

ACTIONS = {"MAP", "NEW", "IGNORE"}


def lit(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def read_rows(path):
    ws = load_workbook(path, data_only=True)["Operators"]
    rows = []
    for r in range(2, ws.max_row + 1):
        op_id = ws.cell(r, 1).value
        if op_id is None or not str(op_id).strip().isdigit():
            break
        rows.append({
            "id": int(op_id),
            "name": ws.cell(r, 2).value,
            "action": (str(ws.cell(r, 5).value or "").strip().upper() or "IGNORE"),
            "person": (str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value else None),
            "new": (str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value else None),
            "researcher": str(ws.cell(r, 8).value or "").strip().upper() == "Y",
            "additional": [p.strip() for p in str(ws.cell(r, 13).value or "").split(",") if p.strip()],
        })
    return rows


def arr(names):
    if not names:
        return "'{}'::text[]"
    return "ARRAY[" + ", ".join(lit(n) for n in names) + "]::text[]"


def convert(path):
    values = []
    for row in read_rows(path):
        a = row["action"]
        if a not in ACTIONS:
            raise ValueError(f"row {row['id']} ({row['name']}): unknown action {a!r}")
        if a == "IGNORE":
            continue
        if a == "MAP" and not row["person"]:
            raise ValueError(f"row {row['id']} ({row['name']}): MAP needs a person")
        if a == "NEW" and not row["new"]:
            raise ValueError(f"row {row['id']} ({row['name']}): NEW needs a name")
        person = row["person"] if a == "MAP" else None
        new = row["new"] if a == "NEW" else None
        researcher = "true" if (a == "NEW" and row["researcher"]) else "false"
        values.append(f"({row['id']}, {lit(a)}, {lit(person)}, {lit(new)}, {researcher}, {arr(row['additional'])})")
    body = ",\n    ".join(values) if values else "(NULL::int, NULL, NULL, NULL, false, '{}'::text[])"
    return f"""-- migrate:up
-- Generated from db/mappings/operator-mapping.xlsx by scripts/operator_mapping_to_sql.py.
-- Links legacy Machine_Operators entries to people: MAP -> existing person,
-- NEW -> create one person per distinct name; that person becomes the main
-- operator and any "additional people" go in the additional-operators
-- junctions. IGNORE rows keep only the operator_name text already on each run.
CREATE TEMP TABLE _opmap (
    machine_operator_id integer PRIMARY KEY,
    action text NOT NULL,
    person_name text,
    new_name text,
    researcher boolean NOT NULL,
    additional text[] NOT NULL
) ON COMMIT DROP;

INSERT INTO _opmap (machine_operator_id, action, person_name, new_name, researcher, additional) VALUES
    {body};
DELETE FROM _opmap WHERE machine_operator_id IS NULL;

INSERT INTO people (full_name, is_operator, is_researcher)
SELECT DISTINCT ON (lower(new_name)) new_name, true, researcher
  FROM _opmap m
 WHERE action = 'NEW'
   AND NOT EXISTS (SELECT 1 FROM people p WHERE lower(p.full_name) = lower(m.new_name))
 ORDER BY lower(new_name), researcher DESC;

DO $$
DECLARE v text;
BEGIN
    SELECT string_agg(DISTINCT n, ', ') INTO v
      FROM (SELECT COALESCE(person_name, new_name) AS n FROM _opmap WHERE action IN ('MAP', 'NEW')
            UNION ALL
            SELECT unnest(additional) FROM _opmap WHERE action IN ('MAP', 'NEW')) names
     WHERE (SELECT count(*) FROM people p WHERE lower(p.full_name) = lower(names.n)) <> 1;
    IF v IS NOT NULL THEN
        RAISE EXCEPTION 'operator mapping: names missing or ambiguous in people: %', v;
    END IF;
END $$;

UPDATE manufacturing_operations o SET operator_person_id = p.person_id
  FROM _opmap m JOIN people p ON lower(p.full_name) = lower(COALESCE(m.person_name, m.new_name))
 WHERE m.action IN ('MAP', 'NEW') AND o.operator = m.machine_operator_id AND o.operator_person_id IS NULL;

UPDATE test_sessions s SET operator_person_id = p.person_id
  FROM _opmap m JOIN people p ON lower(p.full_name) = lower(COALESCE(m.person_name, m.new_name))
 WHERE m.action IN ('MAP', 'NEW') AND s.operator = m.machine_operator_id AND s.operator_person_id IS NULL;

INSERT INTO operation_additional_operators (operation_id, person_id)
SELECT o.operation_id, p.person_id
  FROM _opmap m
  CROSS JOIN LATERAL unnest(m.additional) AS a(name)
  JOIN people p ON lower(p.full_name) = lower(a.name)
  JOIN manufacturing_operations o ON o.operator = m.machine_operator_id
 WHERE m.action IN ('MAP', 'NEW') AND p.person_id IS DISTINCT FROM o.operator_person_id
ON CONFLICT (operation_id, person_id) DO NOTHING;

INSERT INTO session_additional_operators (session_id, person_id)
SELECT s.session_id, p.person_id
  FROM _opmap m
  CROSS JOIN LATERAL unnest(m.additional) AS a(name)
  JOIN people p ON lower(p.full_name) = lower(a.name)
  JOIN test_sessions s ON s.operator = m.machine_operator_id
 WHERE m.action IN ('MAP', 'NEW') AND p.person_id IS DISTINCT FROM s.operator_person_id
ON CONFLICT (session_id, person_id) DO NOTHING;

-- migrate:down
-- Data backfill; not reverted.
SELECT 1;
"""


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    sql = convert(src)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(sql)
    print(f"wrote {out}")
